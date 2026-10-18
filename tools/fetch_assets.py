"""Laedt Vereinslogos und Spielerfotos von Wikipedia/Wikimedia und pflegt data/players.json.

Aufruf (im Repo-Ordner):   python tools/fetch_assets.py
Optionen:                  --min-opp-goals N   Gegner-Torschuetzen erst ab N Toren (Standard 1)
                           --refresh           bereits vorhandene Eintraege neu suchen

Bereits vorhandene Logos/Fotos werden nicht neu geladen. Eigenes Foto: als
assets/players/<vorname-nachname>.jpg ablegen (z. B. jens-stage.jpg) und das Skript laufen lassen. Eintraege in players.json mit
"manual": true werden nie automatisch ueberschrieben -> dort Korrekturen von Hand eintragen.
Benoetigt: pip install openpyxl pillow
"""
import argparse, json, re, sys, time, unicodedata, urllib.error, urllib.parse, urllib.request
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "data" / "werderspiele.xlsx"
CLUBS = ROOT / "data" / "clubs.json"
PLAYERS = ROOT / "data" / "players.json"
LOGOS = ROOT / "assets" / "logos"
PHOTOS = ROOT / "assets" / "players"
UA = "WerderspieleDashboard/1.0 (https://github.com/demuHB/werderspiele; private non-commercial) python-urllib"


def http_get(url):
    for attempt in range(6):
        req = urllib.request.Request(url, headers={"User-Agent": UA})
        try:
            with urllib.request.urlopen(req, timeout=30) as r:
                return r.read()
        except urllib.error.HTTPError as e:
            if e.code != 429 or attempt == 5:
                raise
            time.sleep(5 * (attempt + 1))


def http_json(url):
    return json.loads(http_get(url))


def download(url, dest):
    dest.write_bytes(http_get(url))
    if dest.parent == PHOTOS:
        shrink(dest)
    time.sleep(1)


def shrink(path, width=160):
    """Spielerfotos klein halten (schnelles Laden am Handy). Braucht Pillow, sonst bleibt das Original."""
    try:
        from PIL import Image
    except ImportError:
        return
    im = Image.open(path)
    if im.mode not in ("RGB", "L"):
        bg = Image.new("RGB", im.size, (238, 244, 240))
        im = im.convert("RGBA")
        bg.paste(im, mask=im.split()[3])
        im = bg
    if im.width > width:
        im = im.resize((width, round(im.height * width / im.width)), Image.LANCZOS)
    im.convert("RGB").save(path, "JPEG", quality=82, optimize=True)


def api(lang, **params):
    params.update(format="json", formatversion="2")
    return http_json(f"https://{lang}.wikipedia.org/w/api.php?" + urllib.parse.urlencode(params))


def slug(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def fold(s):
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()


# ---------- Excel lesen (gleiche Regeln wie js/parse.js) ----------
def split_outside_parens(text):
    """An Kommas trennen, aber nicht innerhalb von Klammern ("Ailton (3, FE), Micoud")."""
    parts, depth, cur = [], 0, ""
    for ch in text or "":
        depth += (ch == "(") - (ch == ")")
        if ch == "," and depth == 0:
            parts.append(cur)
            cur = ""
        else:
            cur += ch
    parts.append(cur)
    return parts


def parse_scorers(text):
    """-> [(name, tore, eigentor)]; Klammern wie in js/parse.js: (2), (FE), (3, FE), (3, 2 FE), (ET)"""
    out = []
    for p in split_outside_parens(text):
        p = p.strip()
        if not p or p in ("/", "*"):
            continue
        items = [t.strip() for m in re.findall(r"\(([^)]*)\)", p) for t in m.split(",") if t.strip()]
        n = next((int(t) for t in items if t.isdigit()), None)
        marks = sum(int(m[1] or 1) for t in items if (m := re.fullmatch(r"(\d+)?\s*(FE|HE)", t, re.I)))
        own = any(re.fullmatch(r"(\d+\s*)?ET", t, re.I) for t in items)
        out.append((re.sub(r"\s*\(.*?\)", "", p).strip(), n or max(1, marks), own))
    return out


def collect_scorers():
    """-> Counter[(name, team)] mit team = 'Werder' oder Gegnername (Spieler-Verein, nicht Torseite)."""
    clubs = json.loads(CLUBS.read_text(encoding="utf-8"))["clubs"]
    alias = {a: k for k, v in clubs.items() for a in v.get("aliases", [])}
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["Haupttabelle"]
    cnt = Counter()
    for r in range(5, ws.max_row + 1):
        f, res, sc = ws.cell(r, 6).value, ws.cell(r, 8).value, ws.cell(r, 9).value
        if not f or not res:
            continue
        home, away = [alias.get(x.strip(), x.strip()) for x in f.split(" - ", 1)]
        m = re.match(r"\s*(\d+):(\d+)", str(res))
        h, a = int(m[1]), int(m[2])
        werder_home = home.startswith("Werder")
        opp = away if werder_home else home
        parts = sc.split(" - ") if sc and sc.strip() != "/" else []
        if len(parts) == 1:
            parts = [parts[0], ""] if a == 0 else ["", parts[0]]
        for side, txt in enumerate(parts):
            werder_side = (side == 0) == werder_home
            for name, n, own in parse_scorers(txt):
                team = "Werder" if werder_side != own else opp
                cnt[(name, team)] += n
    return cnt


# ---------- Logos ----------
def fetch_logos():
    clubs = json.loads(CLUBS.read_text(encoding="utf-8"))["clubs"]
    LOGOS.mkdir(parents=True, exist_ok=True)
    for name, c in clubs.items():
        if not c.get("id") or not c.get("wiki"):
            continue
        dest = LOGOS / f"{c['id']}.png"
        if dest.exists():
            continue
        src = None
        for lang in ("en", "de"):
            try:
                # page_image enthaelt auch nicht-freie Wappen (pageimages liefert nur freie)
                pages = api(lang, action="query", titles=c["wiki"], prop="pageprops",
                            redirects=1)["query"]["pages"]
                props = pages[0].get("pageprops", {})
                file = props.get("page_image") or props.get("page_image_free")
                if file:
                    info = api(lang, action="query", titles=f"File:{file}", prop="imageinfo",
                               iiprop="url", iiurlwidth=200)["query"]["pages"][0]
                    src = info.get("imageinfo", [{}])[0].get("thumburl")
            except Exception as e:  # noqa
                print("  Fehler", name, lang, e)
            time.sleep(0.5)
            if src:
                break
        if src:
            download(src, dest)
            print("Logo  ok ", name)
        else:
            print("Logo  FEHLT", name)
        time.sleep(0.3)


# ---------- Spieler ----------
def is_footballer(page):
    cats = " ".join(c["title"] for c in page.get("categories", []))
    return "Fußballspieler" in cats or "Fußballnationalspieler" in cats


def find_player(name, team):
    surname = name.split(". ")[-1]
    hint = "Werder Bremen" if team == "Werder" else team
    res = api("de", action="query", list="search", srsearch=f'{surname} Fußballspieler "{hint}"', srlimit=6)
    titles = [r["title"] for r in res["query"]["search"]]
    if not titles:
        return None
    pages = api("de", action="query", titles="|".join(titles), prop="pageimages|categories|extracts",
                pithumbsize=300, cllimit="max", exintro=1, explaintext=1, exlimit="max", redirects=1)["query"]["pages"]
    by_title = {p["title"]: p for p in pages}
    first = name.split(". ")[0] if ". " in name else None
    for t in titles:
        p = by_title.get(t)
        if not p or not is_footballer(p):
            continue
        tf = fold(t)
        if fold(surname) not in tf and fold(surname) not in fold(p.get("extract", "")[:200]):
            continue
        if first and not re.search(r"\b" + fold(first), tf):
            continue
        return {"full": re.sub(r"\s*\(.*\)$", "", t), "wiki": t,
                "thumb": p.get("thumbnail", {}).get("source")}
    return None


def fetch_manual_photo(entry):
    """Foto fuer einen von Hand gesetzten Wikipedia-Titel laden (de, sonst en)."""
    for lang in ("de", "en"):
        pages = api(lang, action="query", titles=entry["wiki"], prop="pageimages",
                    pithumbsize=300, redirects=1)["query"]["pages"]
        thumb = pages[0].get("thumbnail", {}).get("source")
        if thumb:
            entry["thumb"] = thumb
            entry["img"] = f"assets/players/{slug(entry['full'])}.jpg"
            if not (ROOT / entry["img"]).exists():
                download(thumb, ROOT / entry["img"])
            print("Foto ok   ", entry["wiki"])
            return
        time.sleep(0.5)
    entry["noimg"] = True
    print("Foto fehlt", entry["wiki"])


def commons_thumb(filename, width=300):
    """Thumbnail-URL einer Commons-Datei (z. B. aus Wikidata P18)."""
    pages = http_json("https://commons.wikimedia.org/w/api.php?" + urllib.parse.urlencode(dict(
        action="query", titles=f"File:{filename}", prop="imageinfo", iiprop="url", iiurlwidth=width,
        format="json", formatversion="2")))["query"]["pages"]
    return (pages[0].get("imageinfo") or [{}])[0].get("thumburl")


def find_free_photo(wiki_title):
    """Weitere freie Quellen, wenn die deutsche Wikipedia kein Hauptbild hat:
    1. Wikidata-Bild (P18) des Artikels, 2. Hauptbild der englischen Wikipedia."""
    pages = api("de", action="query", titles=wiki_title, prop="pageprops", redirects=1)["query"]["pages"]
    qid = pages[0].get("pageprops", {}).get("wikibase_item")
    if not qid:
        return None
    ent = http_json("https://www.wikidata.org/w/api.php?" + urllib.parse.urlencode(dict(
        action="wbgetentities", ids=qid, props="claims|sitelinks", sitefilter="enwiki", format="json")))["entities"][qid]
    p18 = ent.get("claims", {}).get("P18")
    if p18:
        thumb = commons_thumb(p18[0]["mainsnak"]["datavalue"]["value"])
        if thumb:
            return thumb, "Wikidata"
    en = ent.get("sitelinks", {}).get("enwiki", {}).get("title")
    if en:
        pg = api("en", action="query", titles=en, prop="pageimages", pithumbsize=300)["query"]["pages"][0]
        thumb = pg.get("thumbnail", {}).get("source")
        if thumb and not pg.get("pageimage", "").lower().endswith((".svg", ".png")):  # Wappen/Logos auslassen
            return thumb, "en.wikipedia"
    return None


def fill_missing_photos(data):
    """Für Spieler mit Wikipedia-Artikel, aber ohne Foto: eigene Datei verwenden oder freie Quellen durchsuchen."""
    for key, e in sorted(data.items()):
        if e.get("img") or not e.get("full"):
            continue
        # 1) Eigenes Foto: assets/players/<vorname-nachname>.jpg|png (z. B. jens-stage.jpg)
        own = next((f for f in (PHOTOS / f"{slug(e['full'])}{ext}" for ext in (".jpg", ".jpeg", ".png", ".webp")) if f.exists()), None)
        if own:
            e.update(img=own.relative_to(ROOT).as_posix(), manual=True, source="eigene Datei")
            print("Foto eigen ", key, "->", e["img"])
            continue
        if not e.get("wiki") or e.get("noimg") == "checked-free":
            continue
        try:
            hit = find_free_photo(e["wiki"])
        except Exception as ex:  # noqa
            print("  Fehler", key, ex)
            continue
        if hit:
            thumb, src = hit
            e.update(thumb=thumb, img=f"assets/players/{slug(e['full'])}.jpg", source=src)
            if not (ROOT / e["img"]).exists():
                download(thumb, ROOT / e["img"])
            print("Foto neu  ", key, "<-", src)
        else:
            e["noimg"] = "checked-free"  # nicht bei jedem Lauf erneut suchen
            print("Foto keins", key)
        time.sleep(0.5)


def fetch_players(min_opp, refresh):
    PHOTOS.mkdir(parents=True, exist_ok=True)
    data = json.loads(PLAYERS.read_text(encoding="utf-8")) if PLAYERS.exists() else {}
    for (name, team), goals in sorted(collect_scorers().items(), key=lambda x: -x[1]):
        key = f"{name}@{team}"
        if team != "Werder" and goals < min_opp:
            continue
        entry = data.get(key)
        if entry and (entry.get("manual") or not refresh):
            if entry.get("img") and not (ROOT / entry["img"]).exists() and entry.get("thumb"):
                download(entry["thumb"], ROOT / entry["img"])
            elif entry.get("manual") and entry.get("wiki") and not entry.get("img") and not entry.get("noimg"):
                fetch_manual_photo(entry)
            continue
        try:
            hit = find_player(name, team)
        except Exception as e:  # noqa
            print("  Fehler", key, e)
            continue
        if not hit:
            data[key] = {"full": name, "wiki": None, "img": None}
            print("Spieler ?  ", key)
        else:
            img = None
            if hit["thumb"]:
                ext = ".png" if hit["thumb"].lower().endswith(".png") else ".jpg"
                img = f"assets/players/{slug(hit['full'])}{ext}"
                if not (ROOT / img).exists():
                    download(hit["thumb"], ROOT / img)
            data[key] = {"full": hit["full"], "wiki": hit["wiki"], "img": img, "thumb": hit["thumb"]}
            print("Spieler ok ", key, "->", hit["wiki"], "(Foto)" if img else "(kein Foto)")
        time.sleep(0.7)
        if len(data) % 20 == 0:
            save_players(data)
    fill_missing_photos(data)
    save_players(data)


def save_players(data):
    PLAYERS.write_text(json.dumps(dict(sorted(data.items())), ensure_ascii=False, indent=1), encoding="utf-8")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser()
    ap.add_argument("--min-opp-goals", type=int, default=1)
    ap.add_argument("--refresh", action="store_true")
    a = ap.parse_args()
    fetch_logos()
    fetch_players(a.min_opp_goals, a.refresh)
